import io
import re
from typing import Tuple

import PyPDF2
from docx import Document


# ---------------------------------------------------------
# Configuration
# ---------------------------------------------------------

MAX_FILE_SIZE_MB = 5
MAX_FILE_SIZE_BYTES = MAX_FILE_SIZE_MB * 1024 * 1024

ALLOWED_EXTENSIONS = {".pdf", ".docx"}

ALLOWED_MIME_TYPES = {
    "application/pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


# ---------------------------------------------------------
# Custom Exceptions
# ---------------------------------------------------------

class ResumeParserError(Exception):
    """Base exception for resume parsing errors."""


class FileValidationError(ResumeParserError):
    """Raised when an uploaded file is invalid."""


class TextExtractionError(ResumeParserError):
    """Raised when text cannot be extracted."""


# ---------------------------------------------------------
# Text Cleaning
# ---------------------------------------------------------

def clean_extracted_text(text: str) -> str:
    """
    Clean extracted resume text while preserving useful structure.
    """

    if not text:
        return ""

    # Normalize line endings
    text = text.replace("\r\n", "\n").replace("\r", "\n")

    # Remove excessive spaces/tabs
    text = re.sub(r"[ \t]+", " ", text)

    # Remove excessive blank lines
    text = re.sub(r"\n{3,}", "\n\n", text)

    # Remove spaces around newlines
    text = re.sub(r" *\n *", "\n", text)

    return text.strip()


# ---------------------------------------------------------
# File Validation
# ---------------------------------------------------------

def validate_file(
    file_data: bytes,
    filename: str,
    content_type: str | None = None,
) -> str:
    """
    Validate file size, extension and basic file signature.

    Returns:
        "pdf" or "docx"
    """

    if not file_data:
        raise FileValidationError(
            "The uploaded file is empty."
        )

    # Size validation
    file_size = len(file_data)

    if file_size > MAX_FILE_SIZE_BYTES:
        size_mb = file_size / (1024 * 1024)

        raise FileValidationError(
            f"File is too large ({size_mb:.2f} MB). "
            f"Maximum allowed size is {MAX_FILE_SIZE_MB} MB."
        )

    # Filename validation
    if not filename:
        raise FileValidationError(
            "Filename is missing."
        )

    extension = "." + filename.lower().split(".")[-1]

    if extension not in ALLOWED_EXTENSIONS:
        raise FileValidationError(
            "Unsupported file type. "
            "Please upload a PDF or DOCX file."
        )

    # Basic content validation
    if extension == ".pdf":

        if not file_data.startswith(b"%PDF"):
            raise FileValidationError(
                "The uploaded file does not appear to be a valid PDF."
            )

        return "pdf"

    if extension == ".docx":

        # DOCX files are ZIP containers and normally start with PK.
        if not file_data.startswith(b"PK"):
            raise FileValidationError(
                "The uploaded file does not appear to be a valid DOCX."
            )

        return "docx"

    raise FileValidationError(
        "Unsupported file format."
    )


# ---------------------------------------------------------
# PDF Extraction
# ---------------------------------------------------------

def extract_text_from_pdf(file_data: bytes) -> str:
    """
    Extract text from a PDF using PyPDF2.
    """

    try:
        reader = PyPDF2.PdfReader(
            io.BytesIO(file_data)
        )

        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise TextExtractionError(
                    "This PDF is password protected."
                )

        extracted_pages = []

        for page in reader.pages:

            page_text = page.extract_text()

            if page_text:
                extracted_pages.append(page_text)

        text = "\n".join(extracted_pages)

        text = clean_extracted_text(text)

        if not text:
            raise TextExtractionError(
                "No selectable text could be extracted from this PDF. "
                "The PDF may contain scanned images."
            )

        return text

    except TextExtractionError:
        raise

    except Exception as error:
        raise TextExtractionError(
            "Failed to read the PDF. "
            "The file may be corrupted or unsupported."
        ) from error


# ---------------------------------------------------------
# DOCX Extraction
# ---------------------------------------------------------

def extract_text_from_docx(file_data: bytes) -> str:
    """
    Extract text from DOCX paragraphs and tables.
    """

    try:
        document = Document(
            io.BytesIO(file_data)
        )

        text_parts = []

        # Normal paragraphs
        for paragraph in document.paragraphs:

            text = paragraph.text.strip()

            if text:
                text_parts.append(text)

        # Tables
        for table in document.tables:

            for row in table.rows:

                row_text = []

                for cell in row.cells:

                    cell_text = cell.text.strip()

                    if cell_text:
                        row_text.append(cell_text)

                if row_text:
                    text_parts.append(" | ".join(row_text))

        text = "\n".join(text_parts)

        text = clean_extracted_text(text)

        if not text:
            raise TextExtractionError(
                "No text could be extracted from this DOCX file."
            )

        return text

    except TextExtractionError:
        raise

    except Exception as error:
        raise TextExtractionError(
            "Failed to read the DOCX file. "
            "The document may be corrupted."
        ) from error


# ---------------------------------------------------------
# Main Parser
# ---------------------------------------------------------

def parse_resume_file(
    file_data: bytes,
    filename: str,
    content_type: str | None = None,
) -> Tuple[str, dict]:
    """
    Validate and extract text from a resume.

    Returns:
        (
            extracted_text,
            metadata
        )
    """

    file_type = validate_file(
        file_data=file_data,
        filename=filename,
        content_type=content_type,
    )

    if file_type == "pdf":

        text = extract_text_from_pdf(
            file_data
        )

    elif file_type == "docx":

        text = extract_text_from_docx(
            file_data
        )

    else:
        raise FileValidationError(
            "Unsupported resume format."
        )

    metadata = {
        "filename": filename,
        "file_type": file_type,
        "file_size_bytes": len(file_data),
        "text_length": len(text),
        "success": True,
    }

    return text, metadata